import io
import re
import urllib.parse
import openpyxl
import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Font

# 🚀 네이버 클라우드 플랫폼 API 인증 정보
NAVER_CLIENT_ID = "bkwdzf72zn"
NAVER_CLIENT_SECRET = "rbzw7TF0T8wDDFB4HqoZ4rNaq14zKDgc8W9ogQX6"

def fix_zipcode(val):
    if pd.isna(val) or val is None:
        return ""
    val_str = str(val).split('.')[0].strip()
    return val_str.zfill(5) if val_str else ""

# --- [ 특정 예외 주소 매핑 사전 ] ---
SPECIAL_EXCEPTIONS = {
    "불로동 268-2": "인천광역시 검단구 금정로 12 (불로동, 신검단중앙역풍경채어바니티)",
    "남산타운": "서울특별시 중구 다산로 32 (신당동, 남산타운)",
    "율곡동 415": "경상북도 김천시 혁신8로 66",
}

def remove_illegal_chars(val):
    if isinstance(val, str):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', val)
    return val

def remove_stray_province(addr_str):
    if not addr_str:
        return addr_str
    tokens = str(addr_str).split()
    new_tokens = []
    seen_prov = False
    provinces = ['강원특별자치도', '강원도', '서울특별시', '부산광역시', '대구광역시', '인천광역시', '광주광역시', '대전광역시', '울산광역시', '세종특별자치시', '경기도', '충청북도', '충청남도', '전라북도', '전라남도', '경상북도', '경상남도', '제주특별자치도']
    for tok in tokens:
        if tok in provinces:
            if not seen_prov:
                seen_prov = True
                new_tokens.append(tok)
            else:
                continue
        else:
            new_tokens.append(tok)
    return ' '.join(new_tokens)

# --- [ 괄호 속 중복 아파트명 제거 및 법정동만 남기는 함수 ] ---
def clean_parentheses(addr_str):
    if not addr_str:
        return addr_str
    s = str(addr_str).strip()
    match = re.search(r'\(\s*([가-힣]+동)\s*,\s*([^)]+)\s*\)', s)
    if match:
        legal_dong = match.group(1).strip()
        s = s.replace(match.group(0), f"( {legal_dong} )")
    return ' '.join(s.split())

# --- [ 상품명 정제 함수 ] ---
def clean_product_name(val):
    if pd.isna(val) or not str(val).strip():
        return val
    s = str(val).strip()
    
    has_film = False
    if ('풀액정필름2장' in s or '추가액정필름' in s) and '필요없음' not in s and ('풀액정' in s or '액정필름' in s):
        has_film = True

    target_texts = [s]
    if '//' in s:
        target_texts.insert(0, s.split('//')[-1])
    if '모델선택:' in s:
        target_texts.insert(0, s.split('모델선택:')[-1])
    if '/' in s:
        target_texts.insert(0, s.split('/')[-1])

    smartphone_match = None
    for t in target_texts:
        smartphone_match = re.search(r'(S\s*10\s*5G|S\s*10\s*플러스|S\s*20\s*플러스|S\s*20\s*울트라|S\s*22\s*플러스|S\s*22\s*울트라|노트\s*10\s*플러스|노트\s*20\s*울트라|노트\s*20\s*플러스|노트\s*\d+\s*울트라|노트\s*\d+\s*플러스|노트\s*\d+|S\s*\d+(?:\s*플러스|\s*울트라|\s*5G|\s*엣지)?|A\s*\d+|Z\s*(?:플립|폴드)\s*\d*|와이드\s*\d+|N\s*\d+|V\s*50S|V\s*50|V\s*\d+|벨벳|윙)', t, re.IGNORECASE)
        if smartphone_match:
            break

    if not smartphone_match:
        smartphone_match = re.search(r'(S\s*10\s*5G|S\s*10\s*플러스|S\s*20\s*플러스|S\s*20\s*울트라|S\s*22\s*플러스|S\s*22\s*울트라|노트\s*10\s*플러스|노트\s*20\s*울트라|노트\s*20\s*플러스|노트\s*\d+\s*울트라|노트\s*\d+\s*플러스|노트\s*\d+|S\s*\d+(?:\s*플러스|\s*울트라|\s*5G|\s*엣지)?|A\s*\d+|Z\s*(?:플립|폴드)\s*\d*|와이드\s*\d+|N\s*\d+|V\s*50S|V\s*50|V\s*\d+|벨벳|윙)', s, re.IGNORECASE)

    if smartphone_match:
        raw_model = smartphone_match.group(1).replace(" ", "").upper()
        if raw_model == 'S105G': target_model = 'S10 5G'
        elif raw_model == 'S10플러스': target_model = 'S10플러스'
        elif raw_model == 'S20플러스': target_model = 'S20플러스'
        elif raw_model == 'S22플러스': target_model = 'S22플러스'
        elif raw_model == '노트10플러스': target_model = '노트10플러스'
        elif raw_model in ['N960', 'SM-N960']: target_model = '노트9'
        elif raw_model == 'V50S': target_model = 'V50S'
        elif raw_model == 'V50': target_model = 'V50'
        elif raw_model == '노트20울트라': target_model = '노트20울트라'
        elif raw_model == '와이드5': target_model = 'A22'
        elif raw_model == '와이드6': target_model = 'A13'
        else: target_model = raw_model

        qty = 1
        if '☞' in s:
            parts = s.split('☞')
            after = parts[1].strip() if len(parts) > 1 else ""
            qty_match = re.match(r'^(\d+)\s*(EA|개)?', after, re.IGNORECASE)
            if qty_match:
                qty = int(qty_match.group(1))
        
        res = f"{target_model}[{qty}]" if qty > 1 else target_model
        if has_film: res += " +필름"
        if '택배' in s or '택배배송' in s: res += " 택배"
        return res
    
    if '케이스' not in s:
        res = s
        if has_film: res += " +필름"
        return res
    
    qty = 1
    if '☞' in s:
        parts = s.split('☞')
        base = parts[0].strip()
        after = parts[1].strip() if len(parts) > 1 else ""
        qty_match = re.match(r'^(\d+)\s*(EA|개)?', after, re.IGNORECASE)
        if qty_match:
            qty = int(qty_match.group(1))
    else:
        base = s
        
    if '모델선택:' in base:
        parts = base.split('모델선택:')
        target = parts[-1].strip()
        if '/' in target: target = target.split('/')[0].strip()
        if ',' in target: target = target.split(',')[0].strip()
    else:
        sub_parts = re.split(r'[/]{1,2}', base)
        target = sub_parts[-1].strip() if len(sub_parts) > 1 else base
        
    target = re.sub(r'\([^)]*\)', '', target).strip()
    target = re.sub(r'케이스', '', target)
    target = re.sub(r'풀액정', '', target)
    target = re.sub(r'2장', '', target)
    target = re.sub(r'클리어', '', target)
    target = re.sub(r'투명', '', target)
    target = re.sub(r'갤럭시', '', target)
    target = ' '.join(target.split())
    
    if target == '와이드5': target = 'A22'
    elif target == '와이드6': target = 'A13'
    
    result_str = f"{target}[{qty}]" if qty > 1 else target
    if '택배' in s or '택배배송' in s: result_str = f"{result_str} 택배"
    if has_film: result_str = f"{result_str} +필름"
        
    return result_str

# --- [ 최종 통합 마스터 주소 변환 엔진 ] ---
def master_juso_converter(keyword, original_zip=""):
    if not keyword or pd.isna(keyword):
        return keyword
        
    kw_str = str(keyword).strip()
    kw_str = clean_parentheses(kw_str)
    
    for target_key, override_addr in SPECIAL_EXCEPTIONS.items():
        if target_key in kw_str:
            return override_addr
            
    trailing_paren = ""
    m_paren = re.search(r'(\([^)]+\))\s*$', kw_str)
    if m_paren:
        trailing_paren = m_paren.group(1).strip()
        kw_str = kw_str[:m_paren.start()].strip()
        
    brands = [
        '아파트', '빌라', '빌', '맨션', '타운', '오피스텔', '주공', '단지',
        '자이', '푸르지오', '힐스테이트', '래미안', '아이파크', '캐슬', 'e편한세상', '편한세상',
        '더샵', '데시앙', '뷰', '써밋', '리슈빌', '하늘채', '포레나', '우미린', '트루엘', '코아루', '휴먼시아'
    ]
    
    cut_idx = len(kw_str)
    for b in brands:
        idx = kw_str.find(b)
        if idx != -1 and idx < cut_idx:
            cut_idx = idx
            
    api_query_part = kw_str[:cut_idx].strip()
    tail_part = kw_str[cut_idx:].strip()
    
    if not api_query_part or len(api_query_part) < 5:
        q = re.sub(r'\s+\d+동\s*\d+호.*$', '', kw_str)
        q = re.sub(r'\s+[가-힣a-zA-Z0-9]+\s*동\s*\d+호.*$', '', q)
        q = re.sub(r'\s+\d+호.*$', '', q)
        api_query_part = q
        tail_part = kw_str[len(q):].strip()

    # 중간 불필요한 구 주소 지번 제거 (예: 명륜동 857 등)
    tail_part = re.sub(r'\s+[가-힣]+동\s+\d+(?:-\d+)?(?=\s+\d+동)', '', tail_part)

    for b in brands:
        if b in tail_part:
            match = re.search(rf'{b}\s+(?:[가-힣]+동\s+)?(\d{{1,4}})-(\d{{1,4}})', tail_part)
            if match:
                d_num = match.group(1)
                h_num = match.group(2)
                sub_pattern = rf'{b}\s+(?:[가-힣]+동\s+)?{d_num}-{h_num}'
                tail_part = re.sub(sub_pattern, rf'{b} {d_num}동 {h_num}호', tail_part)
                
    tail_part = re.sub(r'([가-힣A-Za-z0-9]+아파트)\s+(\d{1,4})-(\d{1,4})', r'\1 \2동 \3호', tail_part)
    tail_part = re.sub(r'\b(\d{1,4})-(\d{1,4})\s*호?\b', r'\1동 \2호', tail_part)

    base_road_addr = api_query_part
    url = f"https://maps.apigw.ntruss.com/map-geocode/v2/geocode?query={urllib.parse.quote(api_query_part)}"
    headers = {
        "X-NCP-APIGW-API-KEY-ID": NAVER_CLIENT_ID,
        "X-NCP-APIGW-API-KEY": NAVER_CLIENT_SECRET
    }
    try:
        response = requests.get(url, headers=headers, timeout=5)
        if response.status_code == 200:
            res = response.json()
            addresses = res.get('addresses', [])
            if addresses:
                addr_info = addresses[0]
                road_addr = addr_info.get('roadAddress', '')
                jibun_addr = addr_info.get('jibunAddress', '')
                base_road_addr = road_addr if road_addr else jibun_addr
    except Exception:
        pass
        
    if not base_road_addr:
        base_road_addr = api_query_part

    parts = [base_road_addr]
    if tail_part:
        parts.append(tail_part)
    if trailing_paren:
        parts.append(trailing_paren)

    final_str = ' '.join(parts)
    final_str = remove_stray_province(final_str)
    final_str = re.sub(r'\(\s*\)', '', final_str)
    return ' '.join(final_str.split())

# --- [ Streamlit 웹 UI 실행 부문 ] ---
st.set_page_config(page_title="스마트샵 주소 변환기", page_icon="🛍️", layout="centered")

st.title("🛍️ 스마트샵 주소 변환 & 엑셀 수정 (네이버 API 연동)")
st.write("요청하신 모든 규칙(정밀 신주소 변환, 브랜드 하이픈 동·호수 정렬, 중복 괄호 정돈)이 적용된 최종 변환기입니다.")

uploaded_file = st.file_uploader("변환할 엑셀 파일(.xlsx, .xls, .xlsb)을 업로드하세요", type=["xlsx", "xls", "xlsb"])

if uploaded_file is not None:
    try:
        df = pd.read_excel(uploaded_file)
    except Exception:
        df = pd.read_excel(uploaded_file, engine='xlrd')

    st.write("### 📄 업로드 데이터 미리보기 (상위 5건)")
    st.dataframe(df.head())

    if 'excel_data' not in st.session_state:
        st.session_state.excel_data = None

    if st.button("🚀 주소 변환 및 서식 적용 시작"):
        with st.spinner("주소를 변환하는 중입니다..."):
            if '우편번호' in df.columns:
                df['우편번호'] = df['우편번호'].apply(fix_zipcode)

            if '선택정보' in df.columns:
                df['선택정보'] = df['선택정보'].apply(clean_product_name)

            if '배송지' in df.columns:
                progress_bar = st.progress(0)
                total = len(df)
                
                converted_addrs = []
                for i, row in df.iterrows():
                    orig_zip = str(row['우편번호']).zfill(5) if '우편번호' in df.columns and not pd.isna(row['우편번호']) else ""
                    res = master_juso_converter(row['배송지'], original_zip=orig_zip)
                    converted_addrs.append(res)
                    progress_bar.progress((i + 1) / total)
                    
                df['배송지'] = converted_addrs
            else:
                st.warning("[주의] '배송지' 컬럼을 찾을 수 없습니다.")

            for col in df.columns:
                df[col] = df[col].apply(remove_illegal_chars)

            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False)
                worksheet = writer.sheets['Sheet1']
                
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    for cell in col:
                        val = str(cell.value or '')
                        length = sum(2 if ord(char) > 127 else 1 for char in val)
                        if length > max_len:
                            max_len = length
                    adjusted_width = max(max_len + 4, 12)
                    worksheet.column_dimensions[col_letter].width = min(adjusted_width, 80)

                font_size_8 = Font(size=8)
                for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=len(df.columns)):
                    worksheet.row_dimensions[row[0].row].height = 18
                    for cell in row:
                        cell.font = font_size_8
                        
                if '우편번호' in df.columns:
                    zip_col_idx = df.columns.get_loc('우편번호') + 1
                    for row in range(2, worksheet.max_row + 1):
                        worksheet.cell(row=row, column=zip_col_idx).number_format = '@'

            st.session_state.excel_data = output.getvalue()
        st.success("🎉 주소 변환 및 상품명 정리가 완료되었습니다!")

    if st.session_state.excel_data is not None:
        st.download_button(
            label="📥 변환된 엑셀 파일 다운로드",
            data=st.session_state.excel_data,
            file_name="우편배송_최종완벽정리_v23.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
